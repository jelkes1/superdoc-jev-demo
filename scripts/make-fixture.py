"""Build the public, entirely fictional 14-page demo fixture. Requires python-docx."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
D=Document(); s=D.sections[0];s.page_width=Inches(8.5);s.page_height=Inches(11)
s.top_margin=s.bottom_margin=Inches(.8);s.left_margin=s.right_margin=Inches(.85)
for name in ['Normal','Body Text']:
 st=D.styles[name];st.font.name='Calibri';st.font.size=Pt(11);st.paragraph_format.space_after=Pt(10);st.paragraph_format.line_spacing=1.16
for name,size in [('Title',28),('Heading 1',18),('Heading 2',12)]:
 st=D.styles[name];st.font.name='Calibri';st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string('000000')
h=s.header.paragraphs[0];h.text='NORTHSTAR CLOUD  /  MERIDIAN LABS';h.style='Caption'
f=s.footer.paragraphs[0];f.text='Fictional demonstration agreement    •    ';field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
sections=[
('Software subscription agreement',[
('','Northstar Cloud, Inc. and Meridian Labs, Inc. enter into this Software Subscription Agreement, effective October 1, 2026. Northstar is the Provider and Meridian is the Customer. This agreement governs the hosted workspace, implementation services and support described in the order form.'),
('Order form','The initial subscription is a twelve-month term for a shared enterprise workspace. The parties will appoint an implementation lead and complete the deployment milestones in Schedule A. Pricing excludes taxes and separately approved professional services.'),
('Contract documents','This agreement, its order form and the two schedules form the complete agreement. An order form may identify quantities and services but does not change the legal terms unless both parties expressly sign the change. References to sections mean sections of this agreement.'),
('Demonstration copy','The companies, commercial terms and transactions in this document are fictional. This sample exists to demonstrate document review and tracked changes; it is not an executed agreement.')]),
('1  Definitions and scope',[
('1.1  Services','Services means the hosted Northstar workspace, documented interfaces, standard support and any implementation work expressly included in an order form. The Services allow authorized users to organize business records and coordinate internal workflows.'),
('1.2  Customer data','Customer Data means files, text and other content submitted by or for Customer to the Services. Customer Data does not include Provider software or documentation. Operational Signals means service telemetry, feature usage and performance measurements collected during operation.'),
('1.3  Authorized users','An Authorized User is an employee or contractor whom Customer permits to access its workspace. Customer will maintain an accurate user list and use individual credentials. Customer remains responsible for instructions given through its account.'),
('1.4  Affiliates','An Affiliate is an entity controlled by, controlling or under common control with a party. Control means ownership of more than half of the voting interests. An affiliate receives access only when named in an order form and subject to the same use restrictions.'),
('1.5  Order of precedence','A signed amendment controls over this agreement only for the subject it expressly addresses. This agreement controls over an inconsistent purchase order. Schedules describe operating requirements and do not alter commercial allocation of risk.')]),
('2  Access and implementation',[
('2.1  Subscription access','During the Subscription Term, Provider grants Customer a nonexclusive right to access the Services for its internal business purposes. Customer may use the documented interfaces to connect its own applications within the limits in the order form.'),
('2.2  Deployment','The parties will agree an implementation schedule within ten business days of the Effective Date. Customer will provide configuration instructions, a test workspace and designated personnel. Provider will identify material dependencies before each milestone begins.'),
('2.3  Acceptance','Customer will evaluate configured workflows against the written acceptance criteria in Schedule A. Customer must identify a reproducible material deviation within ten business days of delivery. Provider will correct the deviation and submit that deliverable for another evaluation.'),
('2.4  Customer responsibilities','Customer is responsible for the accuracy and lawful collection of Customer Data and for selecting appropriate access roles. Customer will not knowingly upload malicious code, defeat access controls, or attempt to extract another tenant’s data.'),
('2.5  Changes to scope','A material change to implementation scope requires a written change order describing deliverables, dates and charges. Neither party is required to start additional work before that change order is signed.')]),
('3  Service operation and support',[
('3.1  Service availability','Provider will use commercially reasonable efforts to meet the availability target in Schedule B. Planned maintenance will ordinarily take place outside Customer’s principal business hours. Provider will notify Customer when planned maintenance may materially interrupt access.'),
('3.2  Support','Customer may submit support requests through the designated support channel. Each request should describe the observed behavior, affected workflow and relevant timestamps. Customer will avoid including passwords or unnecessary personal information in a request.'),
('3.3  Incident coordination','For a material service incident, Provider will assign an incident owner and provide periodic operational updates. The parties will preserve relevant diagnostic information and cooperate to reduce the impact on business operations.'),
('3.4  Product updates','Provider may update the Services while preserving their material contracted functionality. Provider will announce a materially incompatible interface change with reasonable advance notice and provide a documented migration path.'),
('3.5  Exclusions','Service levels exclude outages caused by Customer systems, unauthorized modifications, third-party networks outside Provider’s reasonable control, or suspension permitted under this agreement. The parties will determine the cause using available operational evidence.')]),
('4  Fees and payment',[
('4.1  Subscription fees','Customer will pay the fees stated in the order form. Fees are calculated using the subscribed quantities and the pricing period in that order form. Additional quantities require Customer’s written authorization.'),
('4.2  Payment period','Customer shall pay each undisputed invoice within 45 days after receipt.'),
('4.3  Invoice disputes','Customer must give reasonably detailed notice of an invoice dispute before the applicable due date. The parties will work promptly to resolve the dispute. Customer will pay all undisputed amounts when due.'),
('4.4  Taxes','Fees exclude applicable sales, use and similar transaction taxes. Customer is responsible for those taxes except taxes based on Provider’s net income. Provider will identify separately any tax it is required to collect.'),
('4.5  Expenses','Provider may invoice travel or other out-of-pocket expenses only when Customer approved the category and estimated amount in writing before the expense was incurred. Supporting receipts will accompany the invoice.')]),
('5  Term and renewal',[
('5.1  Initial term','The initial Subscription Term begins on the Effective Date and continues for twelve months. Implementation milestones do not extend or shorten that term unless a signed order form expressly provides otherwise.'),
('5.2  Automatic renewal','The subscription automatically renews for successive twelve-month terms unless either party gives written notice of non-renewal at least 15 days before the end of the then-current term.'),
('5.3  Termination for breach','Either party may terminate this agreement for a material breach that remains uncured thirty days after written notice describing the breach. A notice must identify the contractual obligation at issue and the action reasonably required to cure it.'),
('5.4  Exit assistance','For thirty days after expiration, Provider will make Customer Data available for export through the standard export function. Any additional migration assistance requires a written statement of work. Customer will designate a person to coordinate the export.'),
('5.5  Effect of termination','Termination does not eliminate accrued payment obligations or rights that arose before termination. Confidentiality, ownership and provisions that by their nature should survive remain effective for their stated periods.')]),
('6  Limitation of liability',[
('6.1  Aggregate cap','Each party’s aggregate liability arising out of or relating to this agreement shall not exceed the fees paid or payable under this agreement during the 18 months preceding the event giving rise to the claim.'),
('6.2  Excluded losses','Neither party will be liable for indirect, special or consequential losses, including lost profits or business interruption, to the extent permitted by applicable law. This exclusion applies regardless of the form of action and whether the party was advised of the possibility of such losses.'),
('6.3  Application','The aggregate cap applies across all claims arising under this agreement and its order forms. Multiple claims do not enlarge the aggregate cap. The parties acknowledge that the fees reflect this allocation of commercial risk.'),
('6.4  Mitigation','A party seeking damages will take reasonable steps to mitigate its losses. The parties will cooperate in good faith to preserve evidence relevant to a claim, subject to their confidentiality obligations and applicable law.'),
('6.5  Mandatory rights','Nothing in this agreement excludes liability to the extent such exclusion is prohibited by applicable law. Any remaining limitation continues to apply to the fullest extent permitted.')]),
('7  Customer data and improvement',[
('7.1  Ownership','Customer retains all rights in Customer Data. Customer authorizes Provider to process Customer Data to deliver, secure and support the Services in accordance with this agreement and Customer’s documented instructions.'),
('7.2  Improvement activities','Provider may use Operational Signals to improve its models and services. Such signals may include de-identified excerpts derived from Customer Data where reasonably necessary to evaluate a feature. The improvement program is enabled for enterprise workspaces unless the parties agree otherwise in the implementation plan.'),
('7.3  Processing instructions','Customer’s administrators may configure available processing controls in the workspace. Provider will document which controls apply to each processing activity and will not treat an administrator’s silence as a modification to a signed agreement.'),
('7.4  Subprocessors','Provider may use service providers to support hosting, security and operational delivery. Provider remains responsible for their performance of its contractual obligations and will maintain an up-to-date list of material subprocessors.'),
('7.5  Deletion','Following completion of the agreed export period, Provider will delete Customer Data from active systems in accordance with its documented deletion schedule. Backup copies remain protected and are removed through the normal backup lifecycle.')]),
('8  Confidentiality and security',[
('8.1  Confidential information','Confidential Information includes nonpublic business, technical and financial information disclosed in connection with this agreement. It includes Customer Data, security documentation and the commercial terms of the order form.'),
('8.2  Protection','The receiving party will protect Confidential Information using reasonable care and use it only to perform or exercise rights under this agreement. Access will be limited to personnel and professional advisers with a need to know and appropriate confidentiality obligations.'),
('8.3  Security program','Provider will maintain an information security program appropriate to the Services, including access management, encryption in transit, vulnerability management and incident response procedures. Customer remains responsible for its endpoint devices and user permissions.'),
('8.4  Security incidents','Provider will notify Customer without undue delay after confirming unauthorized access to Customer Data within Provider-controlled systems. The notice will describe the known impact and mitigation steps, with updates as material facts become available.'),
('8.5  Required disclosure','A party may disclose Confidential Information when required by law, provided it gives advance notice where legally permitted and reasonably cooperates with the other party’s efforts to limit the disclosure.')]),
('9  Intellectual property and warranties',[
('9.1  Provider technology','Provider retains all rights in the Services, its software, documentation and independently developed improvements. No right is granted by implication. Customer retains all rights in its pre-existing materials and business processes.'),
('9.2  Feedback','Customer may provide suggestions concerning the Services. Provider may use those suggestions without restriction, but this permission does not authorize disclosure of Customer’s Confidential Information or use of Customer Data outside Section 7.'),
('9.3  Service warranty','Provider warrants that the Services will materially conform to the published documentation during the Subscription Term. Customer will report a claimed nonconformity with sufficient detail for Provider to reproduce it.'),
('9.4  Warranty remedy','Provider will use commercially reasonable efforts to correct a material nonconformity. If Provider cannot do so within a reasonable period, Customer may terminate the affected Service and receive a prorated refund of prepaid unused fees for that Service.'),
('9.5  Authority','Each party represents that it has the authority to enter into and perform this agreement. Neither party makes a promise concerning a specific commercial outcome from use of the Services.')]),
('10  Governing law and disputes',[
('10.1  Governing law','This agreement is governed by the laws of the State of New York, without regard to conflict-of-laws principles.'),
('10.2  Escalation','Before commencing proceedings, a party will give written notice describing the dispute. Each party will appoint a representative with authority to negotiate a resolution. The representatives will meet within fifteen business days of the notice.'),
('10.3  Venue','The state and federal courts located in New York County, New York have exclusive jurisdiction over proceedings arising from this agreement. Each party consents to that jurisdiction and waives an objection based solely on inconvenience of the forum.'),
('10.4  Interim relief','The escalation process does not prevent either party from seeking interim relief needed to protect Confidential Information or intellectual property. Seeking such relief does not waive the obligation to continue good-faith discussions.'),
('10.5  Continued performance','During a dispute, the parties will continue performing undisputed obligations where reasonably practicable. Neither party is required to disclose privileged communications or compromise its legal rights to participate in the escalation process.')]),
('11  General terms',[
('11.1  Notices','Formal notices must be sent to the business contacts in the order form and identify the agreement and the relevant subject. Operational support requests may use the support channel. Each party may update its notice contact by written notice.'),
('11.2  Assignment','Neither party may assign this agreement without the other’s written consent, except in connection with a merger or sale of substantially all related assets where the successor assumes all obligations. An assignment does not relieve accrued obligations.'),
('11.3  Force majeure','A party is excused from delayed performance to the extent caused by events beyond its reasonable control, provided it promptly notifies the other party and takes reasonable steps to resume performance. This provision does not excuse payment for Services already delivered.'),
('11.4  Entire agreement','This agreement supersedes prior discussions concerning its subject. Amendments must be in writing and signed by both parties. A waiver on one occasion does not constitute a continuing waiver.'),
('11.5  Signatures','The parties may execute this agreement in counterparts and by electronic signature. Each counterpart is considered an original and all counterparts together form the same agreement.')]),
('Schedule A  Implementation plan',[
('A.1  Project ownership','Each party will designate an implementation lead. Leads coordinate configuration decisions, access requests and acceptance testing. Changes affecting scope or commercial terms require approval under Section 2.5.'),
('A.2  Working sessions','The parties will hold a weekly working session during implementation. Provider will circulate open actions and Customer will identify a responsible owner for each required decision. Test data will be fictional or otherwise approved for the test environment.'),
('A.3  Acceptance evidence','Acceptance is assessed against the agreed workflows. A successful test must show the configured roles, expected result and any export artifact needed to verify the result. Screenshots supplement the test record but do not replace it.'),
('A.4  Handoff','At completion, Provider will deliver the configuration summary and administrator guidance. Customer will confirm its operational owner and support contacts. Outstanding nonmaterial items will be recorded with agreed completion dates.')]),
('Schedule B  Service levels',[
('B.1  Measurement','Availability is measured monthly for the production workspace using Provider’s monitoring records. The measurement period excludes the events described in Section 3.5. Provider will provide reasonable supporting information for a disputed measurement.'),
('B.2  Priority definitions','Priority reflects the impact on an agreed production workflow. A complete loss of access with no workaround is critical. Degraded functionality with an available workaround is assessed according to its actual business impact.'),
('B.3  Claims','Customer must submit a service-level claim within thirty days after the affected month and identify the dates and workflows involved. The parties will reconcile the monitoring evidence before determining an applicable credit.'),
('B.4  Operational review','The parties may review support trends and recurring incidents at a quarterly service meeting. Any recommendation to change service scope, security commitments or fees requires the separate written approval described in this agreement.')])]
for i,(title,paras) in enumerate(sections):
 if i:D.add_page_break()
 D.add_paragraph(title,'Title' if i==0 else 'Heading 1')
 if i==0:
  t=D.add_table(rows=0,cols=2);t.style='Light Shading Accent 1'
  for a,b in [('Provider','Northstar Cloud, Inc.'),('Customer','Meridian Labs, Inc.'),('Subscription','Enterprise workspace • 120 seats'),('Annual fee','$72,000'),('Payment','As described in Section 4'),('Effective date','October 1, 2026')]:
   c=t.add_row().cells;c[0].text=a;c[1].text=b
 for heading,text in paras:
  if heading:D.add_paragraph(heading,'Heading 2')
  D.add_paragraph(text)
 if i==13:
  t=D.add_table(rows=0,cols=3);t.style='Light Shading Accent 1'
  for row in [('Priority','Initial response','Update frequency'),('Critical','1 hour','Every 2 hours'),('High','4 business hours','Each business day'),('Normal','1 business day','On material change')]:
   for c,v in zip(t.add_row().cells,row):c.text=v
 if i==12:
  t=D.add_table(rows=0,cols=3);t.style='Light Shading Accent 1'
  for row in [('Milestone','Owner','Acceptance evidence'),('Workspace setup','Provider','Configured roles and tenant'),('Workflow trial','Joint','Three documented test runs'),('Administrator handoff','Customer','Named support owner')]:
   for c,v in zip(t.add_row().cells,row):c.text=v
D.core_properties.title='Northstar and Meridian Software Subscription Agreement';D.core_properties.author='SuperDoc demo';D.core_properties.comments='Fictional public demonstration fixture.'
out=Path(__file__).resolve().parents[1]/'public'/'northstar-meridian.docx';D.save(out);print(f'Created {out.name}; {len(sections)} intentional pages')
